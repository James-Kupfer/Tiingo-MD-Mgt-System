"""Tests for src/inventory_tickers.py, run against a synthetic DDC project and workbook."""

import csv
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import inventory_tickers  # noqa: E402
from config import DDC_DIR  # noqa: E402

DDC_CONFIG = """
[symbols]
sources = ["spy", "IWM"]
workbook_path = '{workbook}'
workbook_sheet = "Inventory"
symbol_column = "Symbol"
mapped_symbol_column = "Mapped Symbol"
exclude = ["Excl  Me", "zex"]
"""


@pytest.fixture
def ddc_project(tmp_path: Path) -> Path:
    """A fake DDC root holding a config.toml, the real workbook parser and a workbook."""
    parser = DDC_DIR / "src" / "workbook.py"
    if not parser.is_file():
        pytest.skip(f"DDC project not present at {DDC_DIR}")
    (tmp_path / "src").mkdir()
    shutil.copy(parser, tmp_path / "src" / "workbook.py")

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Inventory"
    sheet.append(["Symbol", "Mapped Symbol", "Status"])
    sheet.append(["aaa", "AAAX", "Own"])  # override wins, upper-cased
    sheet.append(["bbb", None, "Own"])
    sheet.append(["BBB", None, "Own"])  # duplicate of the row above
    sheet.append(
        ["excl me", None, "Own"]
    )  # on the exclusion list, whitespace-insensitive
    sheet.append(["MES CME", None, "Own"])  # broker descriptor, not requestable
    sheet.append(["SPY", None, "Own"])  # also a DDC source

    port = book.create_sheet("Portfolio")
    port.append(["Symbol", "Portfolio%"])
    port.append(["aaa Jan15'27 10 CALL", 0.1])  # leg -> Inventory spelling -> AAAX
    port.append(["NEWCO", 0.1])  # held, no Inventory row yet -> itself
    port.append(["XYZ Mar19'27 5 PUT", 0.1])  # leg of a non-Inventory underlying
    port.append(["FOO LSE", 0.1])  # foreign listing, no Inventory row -> skipped
    port.append(["excl me", 0.1])  # on the exclusion list
    port.append(["ZEX Jan15'27 1 CALL", 0.1])  # leg of an excluded underlying
    port.append([None, None])
    workbook = tmp_path / "book.xlsx"
    book.save(workbook)

    (tmp_path / "config.toml").write_text(
        DDC_CONFIG.format(workbook=workbook), encoding="utf-8"
    )
    return tmp_path


def test_collect_symbols_resolves_filters_and_adds_sources(ddc_project: Path) -> None:
    assert inventory_tickers.collect_symbols(ddc_project) == [
        "AAAX",
        "BBB",
        "IWM",
        "NEWCO",
        "SPY",
        "XYZ",
    ]


def test_collect_symbols_missing_portfolio_sheet_raises(ddc_project: Path) -> None:
    workbook = ddc_project / "book.xlsx"
    book = openpyxl.load_workbook(workbook)
    del book["Portfolio"]
    book.save(workbook)
    with pytest.raises(ValueError, match="Portfolio"):
        inventory_tickers.collect_symbols(ddc_project)


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        ("LIT LSE", "LTHM"),  # whole cell is an Inventory spelling
        ("lit lse", "LTHM"),  # case-insensitive
        ("AAA Jan15'27 1 CALL", "AAAX"),  # leg resolved via Inventory
        ("CCXI WAR Dec'30 11.5 USD CALL", "CCXI"),  # warrant leg, bare underlying
        ("MSFT", "MSFT"),  # single token
        ("KXREACTOR Dec31'26 Grants License CALL (KXREACTOR)", None),
        ("GRX WSE", None),  # foreign, no Inventory row
    ],
)
def test_resolve_portfolio_symbol(cell: str, expected: str | None) -> None:
    lookup = {"lit lse": "LTHM", "aaa": "AAAX"}
    assert (
        inventory_tickers.resolve_portfolio_symbol(cell, lookup, str.casefold)
        == expected
    )


def test_write_ticker_file_uses_tickers_to_update_layout(tmp_path: Path) -> None:
    target = tmp_path / "out.csv"
    inventory_tickers.write_ticker_file(["AAA", "BBB"], target)

    with open(target, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0].keys()) == inventory_tickers.FIELDNAMES
    assert [r["ticker"] for r in rows] == ["AAA", "BBB"]
    assert all(r["startDate"] == "" and r["endDate"] == "" for r in rows)


def test_collect_symbols_missing_workbook_raises(ddc_project: Path) -> None:
    (ddc_project / "book.xlsx").unlink()
    with pytest.raises(FileNotFoundError):
        inventory_tickers.collect_symbols(ddc_project)


def test_main_returns_1_and_writes_nothing_on_failure(
    ddc_project: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (ddc_project / "book.xlsx").unlink()
    target = tmp_path / "Tickers_Inventory.csv"
    monkeypatch.setattr(inventory_tickers, "DDC_DIR", ddc_project)
    monkeypatch.setattr(inventory_tickers, "INVENTORY_TICKERS_FILE", target)
    monkeypatch.setattr(inventory_tickers, "LOG_FILE", tmp_path / "test.log")
    monkeypatch.setattr(inventory_tickers.time, "sleep", lambda _s: None)

    assert inventory_tickers.main() == 1
    assert not target.exists()


def test_retry_recovers_from_transient_read_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls = []
    sleeps = []

    def flaky(_ddc_dir: Path) -> list[str]:
        calls.append(1)
        if len(calls) < 3:
            raise ValueError("workbook mid-save")
        return ["AAA"]

    monkeypatch.setattr(inventory_tickers, "collect_symbols", flaky)
    monkeypatch.setattr(inventory_tickers.time, "sleep", sleeps.append)

    assert inventory_tickers.collect_symbols_with_retry(Path(".")) == ["AAA"]
    assert len(calls) == 3 and len(sleeps) == 2


def test_retry_gives_up_after_configured_attempts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls = []

    def always_locked(_ddc_dir: Path) -> list[str]:
        calls.append(1)
        raise PermissionError("held by Excel")

    monkeypatch.setattr(inventory_tickers, "collect_symbols", always_locked)
    monkeypatch.setattr(inventory_tickers.time, "sleep", lambda _s: None)

    with pytest.raises(PermissionError):
        inventory_tickers.collect_symbols_with_retry(Path("."))
    assert len(calls) == inventory_tickers.INVENTORY_READ_ATTEMPTS


def test_config_errors_are_not_retried(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = []

    def bad_config(_ddc_dir: Path) -> list[str]:
        calls.append(1)
        raise KeyError("symbols")

    monkeypatch.setattr(inventory_tickers, "collect_symbols", bad_config)

    with pytest.raises(KeyError):
        inventory_tickers.collect_symbols_with_retry(Path("."))
    assert len(calls) == 1
